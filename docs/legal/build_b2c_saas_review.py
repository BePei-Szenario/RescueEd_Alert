"""Create a counsel-review copy of the B2C RescueEd Alert terms.

This document is deliberately separate from the currently published/default
legal text in lib/legal-documents.ts. Counsel approval and versioned publication
in the application are separate steps.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "RescueEd_Alert_B2C_SaaS_AGB_Anwaltspruefung.docx"

SECTIONS = [
    (
        "1 Anbieter und Geltungsbereich",
        [
            "Anbieter der digitalen Dienstleistung RescueEd Alert ist RescueEd – Benjamin Peinzger, Mühltor 17, 99986 Niederdorla, Deutschland, E-Mail: info_ra@rescueed.de, Telefon: +49 3601 80 80 191 (nachfolgend „wir“).",
            "Diese Allgemeinen Geschäftsbedingungen gelten für natürliche Personen, die RescueEd Alert überwiegend zu privaten Zwecken über die mobile App nutzen (nachfolgend „Nutzer“). Sie regeln das private App-Konto und die Nutzung der über die App erreichbaren digitalen Dienstleistung. Für Organisationskonten und entgeltliche Einzelbuchungen der Organisations-SaaS gelten gesonderte Bedingungen. Das private App-Abo verschafft keinen Zugang zur Web-SaaS oder zum Unternehmerbereich.",
            "Vorrang haben individuelle Vereinbarungen und zwingende gesetzliche Verbraucherrechte. Bedingungen der App-Stores für den dort abgewickelten Kauf und die Zahlung gelten daneben, soweit sie einschlägig sind."
        ],
    ),
    (
        "2 Registrierung und Zugang",
        [
            "Für die Nutzung ist ein persönliches App-Konto mit zutreffenden Angaben und einer erreichbaren E-Mail-Adresse erforderlich. Das Konto wird nach der vorgesehenen E-Mail-Bestätigung freigeschaltet. Die Registrierung ist für sich genommen kostenlos und löst kein Abonnement aus.",
            "Der Nutzer hält Passwort und Sicherheitscodes geheim und informiert uns unverzüglich über einen vermuteten unbefugten Zugriff. Bei einem konkreten Sicherheits- oder Missbrauchsverdacht können wir den Zugang im erforderlichen Umfang vorübergehend sichern oder einschränken. Wir informieren den betroffenen Nutzer, soweit dies den Schutz des Kontos oder anderer Personen nicht gefährdet."
        ],
    ),
    (
        "3 Leistungsumfang des privaten Abos",
        [
            "Ein aktives, dem App-Konto zugeordnetes Monatsabo ermöglicht das Anlegen neuer Events in der App. Die Zahl der Events ist während der aktiven Abozeit nicht durch ein monatliches Kontingent begrenzt. Ein einzelnes Event darf höchstens 48 Stunden dauern. Für das Anlegen eines im Abo enthaltenen Events wird kein zusätzlicher Eventpreis berechnet.",
            "Zum vereinbarten Funktionsumfang gehören die in der App beschriebenen Möglichkeiten zur Eventverwaltung, Erfassung und Einteilung von Helfern, Verwaltung von Sanitätsmitteln, QR-gestützten Anwesenheitserfassung sowie zur Alarmierung und Bestätigung. Der konkrete Funktionsumfang und die technischen Voraussetzungen werden vor Abschluss des Abos in der App beziehungsweise im Store dargestellt.",
            "Endet das bestätigte Abo, können keine neuen Events angelegt werden. Bereits angelegte Events bleiben grundsätzlich im vorgesehenen Eventzeitraum nutzbar; insbesondere wird ein laufendes Event nicht allein wegen des Aboendes abgebrochen. Gesetzliche Rechte bei Störungen der digitalen Dienstleistung bleiben unberührt."
        ],
    ),
    (
        "4 Abschluss und Preis des Monatsabos",
        [
            "Das private Abo kann ausschließlich als In-App-Abonnement über Google Play oder den Apple App Store abgeschlossen werden. Vor der verbindlichen Bestellung zeigt der jeweilige Store den tatsächlich verlangten Gesamtpreis einschließlich anwendbarer Steuern, die Abrechnungsperiode und die wesentlichen Zahlungsbedingungen an. Für Deutschland ist derzeit ein Preis von 39,00 Euro pro Monat vorgesehen; maßgeblich ist der im Store unmittelbar vor dem Kauf angezeigte Preis.",
            "Die kostenpflichtige Bestellung wird erst durch den vom Nutzer im Store ausdrücklich bestätigten Kauf ausgelöst. Eine bloße Registrierung oder das Öffnen der Aboansicht ist keine kostenpflichtige Bestellung. Die App schaltet neue Events erst frei, nachdem der Store einen aktiven Bezug bestätigt und unser Server ihn dem App-Konto zugeordnet hat. Ein fehlgeschlagener, noch ausstehender oder nicht zuordenbarer Kauf bewirkt keine Freischaltung."
        ],
    ),
    (
        "5 Laufzeit Verlängerung und Kündigung",
        [
            "Das Abo hat eine monatliche Abrechnungsperiode und verlängert sich nach den vor dem Kauf im Store angezeigten Bedingungen, bis es gekündigt wird. Die ordentliche Kündigung kann über die Aboverwaltung des Stores erklärt werden; die App soll einen unmittelbar erreichbaren Weg dorthin anbieten. Die Kündigung verhindert weitere Verlängerungen. Für einen bereits bezahlten Zeitraum richtet sich die Nutzbarkeit nach den maßgeblichen Storebedingungen und zwingendem Recht.",
            "Die Deinstallation der App und die Löschung des RescueEd-Alert-Kontos kündigen das Store-Abo nicht. Vor einer Kontolöschung weisen wir darauf hin, dass das Abo gesondert im jeweiligen Store zu kündigen ist. Das Recht zur außerordentlichen Kündigung aus wichtigem Grund und sonstige gesetzliche Beendigungsrechte bleiben unberührt."
        ],
    ),
    (
        "6 Widerruf und Erstattungen",
        [
            "Verbraucher erhalten vor Abschluss des kostenpflichtigen Abos eine gesonderte Widerrufsbelehrung. Kündigung für künftige Abrechnungsperioden, Widerruf des Vertragsschlusses und Kontolöschung sind unterschiedliche Erklärungen. Der Beginn der Nutzung führt nicht schon für sich genommen zum Erlöschen eines gesetzlichen Widerrufsrechts.",
            "Für Widerruf und Erstattung gelten die gesetzlichen Rechte sowie die für den konkreten Kaufweg einschlägigen Abläufe des jeweiligen Stores. Diese AGB beschränken gesetzliche Ansprüche nicht und ersetzen die gesonderte Widerrufsbelehrung nicht."
        ],
    ),
    (
        "7 Technische Voraussetzungen und Alarmierung",
        [
            "Die App setzt ein kompatibles Mobilgerät, eine unterstützte Betriebssystemversion, eine Datenverbindung und die für einzelne Funktionen erforderlichen Berechtigungen voraus. Wir informieren über wesentliche Kompatibilitätsanforderungen vor dem Aboabschluss und stellen erforderliche Sicherheits- und Funktionsaktualisierungen nach Maßgabe der gesetzlichen Vorschriften bereit.",
            "RescueEd Alert unterstützt die Organisation von Sanitätsdiensten. Die Alarmfunktion ist kein Notrufsystem und darf nicht als alleiniger Kommunikationsweg für zeitkritische Einsätze eingesetzt werden. Benachrichtigungen können sich etwa bei fehlender Verbindung, ausgeschaltetem Gerät, deaktivierten Berechtigungen oder Einschränkungen des Betriebssystems verzögern oder ausbleiben. Der Organisator eines Events muss einen geeigneten unabhängigen Rückfallweg festlegen. Eine Versandanzeige allein bestätigt nicht, dass ein Alarm vom Empfänger wahrgenommen wurde."
        ],
    ),
    (
        "8 Zulässige Nutzung und Inhalte",
        [
            "Der Nutzer darf die Dienstleistung nur rechtmäßig und im vorgesehenen Umfang einsetzen. Er darf keine fremden Konten oder Eventzugänge verwenden und hat QR-Codes, Exporte und sonstige personenbezogene Informationen vor unbefugtem Zugriff zu schützen. Für die Angaben zu Helfern und die Berechtigung zu deren Erfassung ist der jeweilige Eventverantwortliche zuständig.",
            "RescueEd Alert ist nicht für Patientenakten, Diagnosen oder Behandlungsdokumentation vorgesehen. Solche Angaben dürfen nicht in Freitextfelder eingetragen werden. Einzelheiten zur Verarbeitung personenbezogener Daten ergeben sich aus der Datenschutzerklärung."
        ],
    ),
    (
        "9 Bereitstellung Aktualisierungen und Änderungen",
        [
            "Wir stellen die vereinbarten digitalen Funktionen während des maßgeblichen Nutzungszeitraums bereit. Wartung, Sicherheitsmaßnahmen und technische Störungen können die Nutzung zeitweise beeinträchtigen. Über erhebliche planbare Einschränkungen informieren wir, soweit dies möglich und zumutbar ist. Die gesetzlichen Rechte bei unterbliebener Bereitstellung oder Mängeln digitaler Produkte bleiben unberührt.",
            "Änderungen, die über die zur Erhaltung der Vertragsmäßigkeit erforderlichen Aktualisierungen hinausgehen, nehmen wir nur im gesetzlich zulässigen Rahmen vor, etwa aus Sicherheitsgründen oder zur Anpassung an geänderte technische Anforderungen. Soweit eine Änderung die Nutzung mehr als unerheblich beeinträchtigt, informieren wir rechtzeitig auf einem dauerhaften Datenträger über Inhalt, Zeitpunkt und gesetzliche Rechte. Eine Änderung des Aboentgelts richtet sich nach den gesetzlichen Anforderungen und den Regeln des jeweiligen Stores."
        ],
    ),
    (
        "10 Haftung",
        [
            "Wir haften unbeschränkt für Vorsatz und grobe Fahrlässigkeit, für Schäden aus der Verletzung von Leben, Körper oder Gesundheit, nach dem Produkthaftungsgesetz und in anderen Fällen zwingender gesetzlicher Haftung. Bei einfach fahrlässiger Verletzung einer wesentlichen Vertragspflicht ist die Haftung auf den bei Vertragsschluss vorhersehbaren, vertragstypischen Schaden begrenzt. Eine wesentliche Vertragspflicht ist eine Pflicht, deren Erfüllung die ordnungsgemäße Durchführung des Vertrags erst ermöglicht und auf deren Einhaltung der Nutzer regelmäßig vertrauen darf.",
            "Die vorstehenden Regelungen lassen zwingende Verbraucherrechte und eine ausdrücklich übernommene Garantie unberührt. Die fachliche Verantwortung für medizinische Entscheidungen und die Einsatzleitung wird nicht auf die App übertragen."
        ],
    ),
    (
        "11 Geltende Fassung und Schlussbestimmungen",
        [
            "Für den Abschluss des Abos gilt die vor dem Kauf zugänglich gemachte Fassung dieser AGB. Neue Fassungen erhalten eine Versionsnummer. Eine Änderung laufender Verträge erfolgt nur nach Maßgabe einer wirksamen Vereinbarung oder einer gesetzlichen Grundlage; eine bloße Veröffentlichung einer neuen Fassung genügt hierfür nicht.",
            "Es gilt deutsches Recht. Zwingende Verbraucherschutzvorschriften des Staates, in dem der Nutzer seinen gewöhnlichen Aufenthalt hat, bleiben unberührt. Für Fragen zu Konto und Dienstleistung ist RescueEd – Benjamin Peinzger unter info_ra@rescueed.de erreichbar."
        ],
    ),
]

REVIEW_NOTES = [
    "Vertragspartner des kostenpflichtigen Store-Abos, Rechnungsstellung, Widerrufsadressat und Erstattungsweg anhand der konkreten Google- und Apple-Verträge prüfen und in AGB, Widerrufsbelehrung sowie Kaufansicht einheitlich abbilden.",
    "Einordnung des fortlaufenden App-Zugangs als digitale Dienstleistung, Wirksamkeit der gesonderten Widerrufsbelehrung und möglicher Wertersatz bei sofortigem Leistungsbeginn prüfen.",
    "Prüfen, ob die Store-Kaufansicht, AGB-Einbeziehung und dauerhaft speicherbare Vertragsbestätigung sämtliche Verbraucherinformationen enthalten. Ebenso den tatsächlich vorhandenen Kündigungsweg in der App prüfen.",
    "Produktangaben mit der veröffentlichten App abgleichen: 39,00 Euro monatlich in Deutschland, 48 Stunden je Event, unbegrenzte Zahl neu angelegter Events und Fortführung bestehender Events nach Aboende.",
    "Datenschutzrollen bei vom privaten Eventverantwortlichen angelegten Helfern, Löschfristen, Exportmöglichkeiten und Umgang mit medizinischen Freitexten gesondert prüfen.",
]


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.widow_control = True


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.84)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(7)

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True

    heading = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(5)
    heading.paragraph_format.keep_with_next = True

    document.add_paragraph("Allgemeine Geschäftsbedingungen für das private RescueEd Alert App Abo", "Title")
    meta = document.add_paragraph("Prüffassung für anwaltliche Prüfung | Version 2026-09-18-b2c-1 | 18. September 2026")
    meta.paragraph_format.space_after = Pt(13)
    add_body(document, "Diese Bedingungen betreffen ausschließlich das private Monatsabo in der RescueEd Alert App. Ein Zugang zur Web-SaaS für Organisationen ist damit nicht verbunden.")

    for heading_text, paragraphs in SECTIONS:
        document.add_paragraph(heading_text, "Heading 1")
        for text in paragraphs:
            add_body(document, text)

    document.add_page_break()
    document.add_paragraph("Prüfpunkte für die anwaltliche Freigabe", "Title")
    add_body(document, "Der folgende Abschnitt dient ausschließlich der Prüfung. Er ist nicht Bestandteil der gegenüber Verbrauchern zu veröffentlichenden AGB.")
    for item in REVIEW_NOTES:
        document.add_paragraph(item, style="List Bullet")

    document.add_paragraph("Rechtsgrundlagen für die Prüfung", "Heading 1")
    add_body(document, "Insbesondere BGB §§ 305 ff., 312 ff., 327 ff., 355 und 356 sowie die im Zeitpunkt der Veröffentlichung gültigen Store-Vertrags- und Abonnementbedingungen von Google und Apple.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("RescueEd Alert | B2C Prüffassung")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(80, 80, 80)

    document.core_properties.title = "Allgemeine Geschäftsbedingungen für das private RescueEd Alert App Abo"
    document.core_properties.subject = "B2C App Abonnement zur anwaltlichen Prüfung"
    document.core_properties.author = "RescueEd"
    document.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
