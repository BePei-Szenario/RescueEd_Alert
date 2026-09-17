from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = "docs/legal/RescueEd_Alert_AGB_Prueffassung.docx"

sections = [
    ("Vertragspartner und Geltungsbereich", [
        "Anbieter von RescueEd Alert ist RescueEd – Benjamin Peinzger, Mühltor 17, 99986 Niederdorla, Deutschland (nachfolgend „RescueEd“). Diese Allgemeinen Geschäftsbedingungen gelten für die Bereitstellung der Webanwendung RescueEd Alert und der zugehörigen mobilen App an Organisationen, Unternehmer und Privatpersonen (nachfolgend „Kunde“). Verbraucher im Sinne dieser Bedingungen sind natürliche Personen, die den Vertrag überwiegend zu Zwecken schließen, die weder ihrer gewerblichen noch ihrer selbstständigen beruflichen Tätigkeit zugerechnet werden können.",
        "Abweichende Bedingungen des Kunden gelten nur, wenn RescueEd ihnen ausdrücklich zustimmt. Individuelle Vereinbarungen haben Vorrang. Die konkrete Bestellung bestimmt das gebuchte Event, die dort angezeigte Helferzahl, Laufzeit und den Preis. Ein gesondert vereinbartes Service Level Agreement regelt ausschließlich die darin bezeichneten Leistungs- und Supportparameter; die Auftragsverarbeitung personenbezogener Daten wird in der gesonderten AVV geregelt.",
    ]),
    ("Registrierung und Kundenkonto", [
        "Für eine Organisation registriert sich eine hierzu befugte Ansprechperson; eine Privatperson kann sich für die eigene Nutzung registrieren. Der Kunde gibt vollständige und zutreffende Kontakt- und Rechnungsdaten an und hält sie aktuell. Die Registrierung allein löst keine kostenpflichtige Eventbuchung aus. Ein Kundenkonto wird nach Bestätigung der erforderlichen Vertragsunterlagen, Prüfung der E-Mail-Adresse und Festlegung eines Passworts über den zugesandten Link eingerichtet.",
        "Der Kunde verwaltet die Zugriffsberechtigungen seiner Nutzer und stellt sicher, dass Zugangsdaten, Sicherheitscodes und Einladungslinks nicht an unbefugte Personen gelangen. Er informiert RescueEd unverzüglich über einen vermuteten Missbrauch. RescueEd darf einen Zugang bei konkreten Anhaltspunkten für Missbrauch oder erhebliche Sicherheitsrisiken vorübergehend sperren; soweit möglich wird der Kunde hierüber und über die Wiederherstellung informiert.",
    ]),
    ("Bestellung und Vertragsschluss für Events", [
        "Unmittelbar vor Abgabe einer kostenpflichtigen Bestellung zeigt RescueEd Alert die wesentlichen Eventdaten, den konkreten Gesamtpreis einschließlich Umsatzsteuer und die Rechnungsdaten an. Der Kunde kann seine Eingaben vor der abschließenden Bestätigung prüfen und berichtigen. Mit dem Klick auf die Schaltfläche „zahlungspflichtig bestellen“ bucht der Kunde das angezeigte Event verbindlich; damit kommt der Vertrag über dieses Event zustande und das Entgelt entsteht. RescueEd schaltet das Event unmittelbar frei und bestätigt die Buchung per E-Mail auf einem dauerhaften Datenträger. Scheitert die Eventanlage oder Freischaltung, wird die Buchung nicht als erfolgreich abgeschlossen behandelt und es entsteht kein Entgelt.",
        "Jedes kostenpflichtige Event ist eine gesonderte Buchung. Ohne erneute ausdrückliche Buchung entstehen keine weiteren Evententgelte. Eine Zahlung wird nicht automatisch von einem Konto oder einer Karte eingezogen. Kann RescueEd eine Bestellung aus technischen oder rechtlichen Gründen nicht freischalten, entsteht für dieses Event kein Vergütungsanspruch.",
    ]),
    ("Leistungsumfang und Nutzungsgrenzen", [
        "RescueEd Alert unterstützt die Organisation von Sanitätsdiensten durch Eventanlage, QR-basierten Check-in und Check-out, Anwesenheitsübersicht, Einteilung von Helfern auf Sanitätsmittel, Alarmierung zugeteilter Teams sowie die im Produkt verfügbaren Listen und Exporte. Der konkrete Funktionsumfang ergibt sich aus der bei der Bestellung angezeigten Leistungsbeschreibung. Die Zahl der voraussichtlichen Helfer bestimmt das gewählte Event-Paket und dessen Preis.",
        "Ein reguläres Event darf höchstens 48 Stunden dauern. Eine längere Laufzeit ist nur möglich, wenn RescueEd für die Organisation ausdrücklich die Option „Dauernutzer“ freigeschaltet hat. Die Freischaltung dieser Option und eine etwaige kostenlose Nutzung sind im Kundenkonto erkennbar. Für ein als kostenlos angelegtes Event entsteht kein Evententgelt; bereits zuvor kostenpflichtig gebuchte Events werden dadurch nicht rückwirkend kostenfrei.",
        "Die mobile App benötigt ein kompatibles Endgerät, eine passende Betriebssystemversion, die erforderlichen Berechtigungen und eine funktionierende Datenverbindung. Für die Verbindung mit einem Event kann ein berechtigter Helfer einen QR-Code nutzen. Der dadurch gewährte Zugang ist auf das jeweilige Event und die vorgesehenen Helferfunktionen beschränkt und endet spätestens mit dem Auschecken oder dem Ende des Events. Der Kunde ist dafür verantwortlich, QR-Codes nur an berechtigte Personen weiterzugeben.",
    ]),
    ("Alarmierung und operative Verantwortung", [
        "RescueEd Alert ist ein unterstützendes Organisations- und Benachrichtigungssystem. Es ist weder ein Notrufsystem noch ein medizinisches Dokumentationssystem. Die Zustellung oder Wahrnehmung einer Alarmierung kann insbesondere durch fehlende Internetverbindung, Geräteeinstellungen, Energiesparmechanismen, Betriebssysteme, Push-Dienste oder technische Störungen verzögert werden oder ausbleiben. Auch eine in der Anwendung angezeigte technische Versendung belegt nicht, dass ein Empfänger den Alarm wahrgenommen hat.",
        "Der Kunde plant und betreibt deshalb einen eigenständigen, für den Sanitätsdienst geeigneten primären Alarm- und Kommunikationsweg sowie angemessene Rückfallebenen. Er kontrolliert eigenverantwortlich Besetzung, Verfügbarkeit und Rückmeldungen der Sanitätsmittel. Medizinische Entscheidungen, Einsatzleitung und die Erfüllung öffentlich-rechtlicher oder berufsrechtlicher Pflichten verbleiben beim Kunden und seinen Einsatzkräften. Diese Regelung beschränkt keine zwingenden gesetzlichen Ansprüche wegen eines von RescueEd zu vertretenden Fehlers.",
    ]),
    ("Pflichten des Kunden und zulässige Inhalte", [
        "Der Kunde darf RescueEd Alert nur rechtmäßig und im vereinbarten Umfang nutzen. Er legt Helfer und Sanitätsmittel nur mit zutreffenden Angaben an, vergibt Berechtigungen sparsam und entfernt nicht mehr berechtigte Personen zeitnah. Er schützt Ausdrucke, PDFs und QR-Codes mit personenbezogenen oder sicherheitsrelevanten Daten vor unbefugtem Zugriff und nutzt die angebotenen Export- und Löschfunktionen verantwortungsvoll.",
        "Die Anwendung ist nicht für Patientenakten, Diagnosen, Behandlungsdokumentation oder andere planmäßige Verarbeitung besonderer Kategorien personenbezogener Daten bestimmt. Solche Informationen dürfen insbesondere nicht in Freitextmeldungen eingetragen werden. Der Kunde informiert seine Nutzer darüber und sorgt für eine geeignete Rechtsgrundlage und Information der betroffenen Helfer. Sicherheitsvorfälle oder missbräuchliche Inhalte teilt er RescueEd unverzüglich mit.",
    ]),
    ("Preise Rechnungsstellung und Zahlung", [
        "Für ein Event gilt ausschließlich der unmittelbar vor der verbindlichen Bestellung angezeigte Preis. Die Preisansicht weist den Bruttopreis, den enthaltenen Umsatzsteuerbetrag und den Nettobetrag aus. Preisänderungen wirken nicht auf bereits freigeschaltete Events zurück. Individuell freigeschaltete kostenlose Nutzung wird bei der jeweiligen Bestellung mit 0 Euro angezeigt.",
        "Das Entgelt für eine kostenpflichtige Eventbuchung entsteht mit dem erfolgreichen Klick auf „zahlungspflichtig bestellen“ und nicht erst mit der späteren Rechnung. Kostenpflichtige Buchungen werden am Monatsende gesammelt in Rechnung gestellt. RescueEd sendet die Rechnung an die vom Kunden hinterlegte Rechnungsadresse und Rechnungs-E-Mail-Adresse. Eine automatische Abbuchung findet nicht statt. Die Fälligkeit ergibt sich aus der Rechnung; gesetzliche Regeln zum Zahlungsverzug bleiben unberührt. Der Kunde prüft seine Rechnungsdaten vor jeder Bestellung und teilt Änderungen unverzüglich mit.",
    ]),
    ("Eventende Stornierung und Kontolöschung", [
        "Ein Event endet mit der vereinbarten Laufzeit oder durch eine frühere Beendigung in der Anwendung. Eine technische Löschung entfernt den laufenden Zugriff auf das Event und löst, soweit verfügbar, den Export der Eventdetails aus. Sie ist nicht ohne Weiteres eine Stornierung der kostenpflichtigen Buchung. Ob und in welcher Höhe ein Entgelt wegen einer vorzeitigen Beendigung entfällt oder erstattet wird, richtet sich nach einer individuellen Stornovereinbarung sowie den gesetzlichen Ansprüchen des Kunden, insbesondere bei nicht oder mangelhaft erbrachter Leistung.",
        "Der Kunde kann sein Konto nach dem vorgesehenen Bestätigungsverfahren löschen. Die Löschung beendet den Zugang und laufende Events. Bereits entstandene Zahlungsansprüche und gesetzliche Aufbewahrungspflichten bleiben unberührt. Erforderliche Vertrags-, Rechnungs- und Rechtsnachweise werden getrennt und nur im rechtlich zulässigen Umfang weiter aufbewahrt; operative Daten werden nach der AVV und der Datenschutzerklärung behandelt.",
    ]),
    ("Widerrufsrecht für Verbraucher", [
        "Verbrauchern kann bei einer online gebuchten Eventleistung ein gesetzliches Widerrufsrecht zustehen. Über Voraussetzungen, Frist, Ausübung und Folgen des Widerrufs informiert RescueEd Verbraucher gesondert vor Abgabe der Bestellung und in der Buchungsbestätigung. Die Zahlungspflicht bei Bestellung, die sofortige Freischaltung und eine spätere Monatsrechnung beseitigen das Widerrufsrecht nicht automatisch. Zwingende Verbraucherrechte, insbesondere bei mangelhafter Leistung, bleiben unberührt.",
        "Soll die Leistung auf ausdrücklichen Wunsch des Verbrauchers schon vor Ablauf der Widerrufsfrist beginnen, werden dafür erforderliche Erklärungen und Informationen gesondert im Bestellprozess eingeholt beziehungsweise bereitgestellt. Ein möglicher Wertersatz oder ein Erlöschen des Widerrufsrechts richtet sich ausschließlich nach den gesetzlichen Voraussetzungen. Diese AGB ersetzen keine gesonderte Widerrufsbelehrung.",
    ]),
    ("Bereitstellung Wartung und Support", [
        "RescueEd stellt die vertraglich geschuldeten Funktionen während der gebuchten Eventlaufzeit bereit. Geplante Wartung, Sicherheitsupdates und Störungsbehebung können die Nutzung zeitweise beeinträchtigen. RescueEd informiert über absehbare erhebliche Einschränkungen, soweit dies praktisch möglich ist. Verbindliche Verfügbarkeitswerte, Servicezeiten und Reaktionsziele gelten nur, soweit sie in einem wirksam vereinbarten und ausgefüllten SLA festgelegt sind.",
        "Der Kunde meldet Funktionsstörungen mit Zeitpunkt, betroffener Funktion und nachvollziehbarer Beschreibung an den vereinbarten Kontakt. Er übermittelt dabei keine Patientendaten. RescueEd prüft gemeldete Störungen und bemüht sich um eine angemessene Behebung; gesetzliche Mängelrechte bleiben unberührt.",
    ]),
    ("Datenschutz und Vertraulichkeit", [
        "Für Konto-, Vertrags-, Abrechnungs- und Sicherheitsdaten, die RescueEd für eigene Zwecke verarbeitet, gilt die Datenschutzerklärung. Soweit RescueEd Event-, Helfer- und Alarmdaten im Auftrag des Kunden verarbeitet, gelten die gesonderte AVV und die darin beschriebenen Weisungen, technischen Maßnahmen und Löschprozesse. Die bloße Kenntnisnahme der Datenschutzerklärung ist keine datenschutzrechtliche Einwilligung.",
        "Die Parteien behandeln nicht öffentlich bekannte Geschäfts- und Betriebsinformationen der jeweils anderen Partei vertraulich. Gesetzliche Offenlegungspflichten und die Verarbeitung, die zur Vertragserfüllung oder Rechtsverfolgung erforderlich ist, bleiben unberührt. Nach Vertragsende bestehen Vertraulichkeitspflichten fort, solange ein berechtigtes Geheimhaltungsinteresse besteht.",
    ]),
    ("Nutzungsrechte", [
        "RescueEd räumt dem Kunden für die Vertragsdauer ein einfaches, nicht übertragbares Recht ein, die Webanwendung und App im vertraglich vereinbarten Umfang für eigene Organisationszwecke zu nutzen. Der Kunde darf die Software nicht als eigenen Dienst an Dritte weiterverkaufen oder unbefugten Dritten Zugang verschaffen. Zwingende gesetzliche Rechte, insbesondere zur Interoperabilität, bleiben unberührt. An vom Kunden eingebrachten Daten erhält RescueEd nur die zur Vertragserfüllung und nach der AVV erforderlichen Nutzungsbefugnisse.",
    ]),
    ("Haftung", [
        "RescueEd haftet unbeschränkt bei Vorsatz und grober Fahrlässigkeit, bei Schäden aus der Verletzung von Leben, Körper oder Gesundheit sowie nach zwingenden gesetzlichen Vorschriften. Bei einfacher Fahrlässigkeit haftet RescueEd für die Verletzung wesentlicher Vertragspflichten, deren Erfüllung die Durchführung des Vertrags erst ermöglicht und auf deren Einhaltung der Kunde regelmäßig vertrauen darf; in diesem Fall ist die Haftung auf den bei Vertragsschluss vorhersehbaren, vertragstypischen Schaden begrenzt. Im Übrigen ist die Haftung für einfache Fahrlässigkeit ausgeschlossen, soweit gesetzlich zulässig.",
        "Die vorstehenden Regeln gelten auch für gesetzliche Vertreter und Erfüllungsgehilfen von RescueEd. Eine ausdrücklich übernommene Garantie bleibt unberührt. Die Pflicht des Kunden, zumutbare Maßnahmen zur Schadensminderung und zur Sicherung seiner eigenen Daten und Einsatzkommunikation zu ergreifen, bleibt bestehen; dies führt nicht zu einer pauschalen Freistellung von RescueEd.",
    ]),
    ("Änderungen der Bedingungen", [
        "Für eine Eventbuchung gilt die Fassung dieser AGB, die der Kunde bei Abgabe der Bestellung bestätigt hat. Neue Fassungen werden mit Versionsnummer bereitgestellt und gelten für künftige Buchungen erst nach erneuter ausdrücklicher Bestätigung. Bereits geschlossene Eventverträge werden durch eine neue Fassung nicht einseitig geändert. Individuelle Vereinbarungen bedürfen der Bestätigung beider Parteien.",
    ]),
    ("Schlussbestimmungen", [
        "Es gilt deutsches Recht. Zwingende gesetzliche Schutzvorschriften bleiben unberührt. Ist eine Bestimmung dieser AGB unwirksam, richtet sich der Vertrag insoweit nach den gesetzlichen Vorschriften; die Wirksamkeit der übrigen Bestimmungen bleibt unberührt.",
        "Vertragliche Mitteilungen können an die im Kundenkonto hinterlegte E-Mail-Adresse beziehungsweise an den von RescueEd benannten Kontakt erfolgen. Der Kunde hält seine Kontaktadresse empfangsbereit und aktuell. Für Fragen zum Vertrag: info_ra@rescueed.de; Telefon +49 3601 80 80 191.",
    ]),
]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.84), Inches(.73)
sec.left_margin, sec.right_margin = Inches(.86), Inches(.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name, normal.font.size, normal.font.color.rgb = "Arial", Pt(10.5), RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for style_name, size, before, after in [("Title", 18, 0, 12), ("Heading 1", 12.5, 15, 6)]:
    style = styles[style_name]
    style.font.name, style.font.size, style.font.color.rgb = "Arial", Pt(size), RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

doc.add_paragraph("Allgemeine Geschäftsbedingungen RescueEd Alert", "Title")
meta = doc.add_paragraph()
meta.add_run("Prüffassung  |  Stand 17. September 2026  |  SaaS und mobile App")
meta.paragraph_format.space_after = Pt(12)
intro = doc.add_paragraph("Diese Bedingungen regeln das Kundenkonto und die Buchung einzelner Events in RescueEd Alert für Organisationen und Privatpersonen. Maßgeblich für eine kostenpflichtige Buchung sind die im Bestellschritt angezeigten Eventdaten und Preise. Mit dem Klick auf „zahlungspflichtig bestellen“ entsteht die Zahlungspflicht; Verbraucherrechte bleiben unberührt. Alarmierungen über die App unterstützen den Sanitätsdienst, ersetzen aber keinen eigenständigen Primäralarmweg.")
intro.paragraph_format.space_after = Pt(11)

for heading, paragraphs in sections:
    doc.add_paragraph(heading, "Heading 1")
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.widow_control = True

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("RescueEd Alert  ·  AGB Prüffassung  ·  17. September 2026")
run.font.name, run.font.size, run.font.color.rgb = "Arial", Pt(8), RGBColor(85, 85, 85)

doc.core_properties.title = "Allgemeine Geschäftsbedingungen RescueEd Alert"
doc.core_properties.subject = "Organisationskonto und Eventbuchungen"
doc.core_properties.author = "RescueEd"
doc.save(OUTPUT)
print(OUTPUT)
