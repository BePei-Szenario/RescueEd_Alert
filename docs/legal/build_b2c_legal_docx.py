"""Build the B2C review copies from the same text used by the app defaults."""

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (ROOT / "lib" / "legal-documents.ts").read_text(encoding="utf-8")
OUT_DIR = ROOT / "docs" / "legal"


def content_for(key: str) -> tuple[str, str, str]:
    pattern = re.escape(key) + r':\{title:"([^"]+)",version:"([^"]+)",content:`(.*?)`\}'
    match = re.search(pattern, SOURCE, re.S)
    if not match:
        raise ValueError(f"Missing legal default: {key}")
    return match.group(1), match.group(2), match.group(3)


def build(key: str, file_name: str, word_title: str) -> None:
    _, version, content = content_for(key)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.82)
    section.bottom_margin = Inches(.75)
    section.left_margin = Inches(.88)
    section.right_margin = Inches(.84)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)

    for name, size, before, after in (("Title", 18, 0, 11), ("Heading 1", 12, 14, 6)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    doc.add_paragraph(word_title, "Title")
    meta = doc.add_paragraph(f"Prüffassung  |  Version {version}  |  Stand 17. September 2026")
    meta.paragraph_format.space_after = Pt(13)
    for block in content.split("\n\n"):
        lines = block.split("\n")
        if len(lines) == 1 and (re.match(r"^\d+\. ", block) or block in {
            "Widerrufsrecht", "Folgen des Widerrufs", "Muster-Widerrufsformular", "Hinweis zur Abgrenzung"
        }):
            doc.add_paragraph(block, "Heading 1")
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.widow_control = True
            for index, line in enumerate(lines):
                if index:
                    paragraph.add_run().add_break()
                paragraph.add_run(line)

    doc.core_properties.title = word_title
    doc.core_properties.subject = "B2C App Abonnement RescueEd Alert"
    doc.core_properties.author = "RescueEd"
    target = OUT_DIR / file_name
    doc.save(target)
    print(target)


build("agb_b2c", "RescueEd_Alert_B2C_AGB_Prueffassung.docx", "Allgemeine Geschäftsbedingungen für private App Abonnements")
build("widerruf", "RescueEd_Alert_B2C_Widerruf_Prueffassung.docx", "Widerrufsbelehrung für das private App Abo")
